"""Renders an EmployeeProfile (app/services/graph/queries.py — pure graph
data, never LLM-generated) into DOCX and PDF CVs (spec section 26).

Deliberately does NOT use the LLM to write a "professional summary" or any
other prose about the employee. Every line in the output comes directly
from a graph relationship that was actually retrieved — the same
citation-safety principle as Phase 9's answer generation, but simpler here
since there's no natural-language synthesis step to get wrong: rendering
structured data into a template cannot hallucinate. A traceability footer
(pg_id + generation timestamp) is included on every CV so it's always
clear which employee record and when.
"""

import io
from datetime import UTC, datetime

from docx import Document as DocxDocument
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.services.graph.queries import EmployeeProfile


def _skill_line(skill) -> str:
    line = skill.name
    extras = []
    if skill.years_experience:
        extras.append(f"{skill.years_experience}y")
    if skill.proficiency:
        extras.append(skill.proficiency)
    if extras:
        line += f" ({', '.join(extras)})"
    return line


def _certification_line(cert) -> str:
    line = cert.name
    if cert.expiry_date:
        line += f" (expires {cert.expiry_date})"
    return line


def _project_line(project) -> str:
    parts = [project.name]
    if project.role_on_project:
        parts.append(f"— {project.role_on_project}")
    context = [p for p in (project.industry, project.client) if p]
    if context:
        parts.append(f"({', '.join(context)})")
    return " ".join(parts)


def _traceability_footer(profile: EmployeeProfile) -> str:
    timestamp = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    return (
        f"Generated from verified employee records (employee ref: {profile.pg_id}) "
        f"on {timestamp}. All information traceable to source HR/project records; "
        f"not AI-generated prose."
    )


def render_cv_docx(profile: EmployeeProfile) -> bytes:
    doc = DocxDocument()

    doc.add_heading(profile.full_name, level=0)
    subtitle_parts = [p for p in (profile.department_name, profile.company_name) if p]
    if subtitle_parts:
        doc.add_paragraph(" — ".join(subtitle_parts))
    doc.add_paragraph(f"{profile.years_experience} years of professional experience")
    if profile.roles:
        doc.add_paragraph(f"Current role: {profile.roles[0].title}")

    if profile.skills:
        doc.add_heading("Skills", level=1)
        for skill in profile.skills:
            doc.add_paragraph(_skill_line(skill), style="List Bullet")

    if profile.certifications:
        doc.add_heading("Certifications", level=1)
        for cert in profile.certifications:
            doc.add_paragraph(_certification_line(cert), style="List Bullet")

    if profile.projects:
        doc.add_heading("Project Experience", level=1)
        for project in profile.projects:
            doc.add_paragraph(_project_line(project), style="List Bullet")

    doc.add_paragraph()
    footer = doc.add_paragraph(_traceability_footer(profile))
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_cv_pdf(profile: EmployeeProfile) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=LETTER, topMargin=0.75 * inch, bottomMargin=0.75 * inch
    )
    styles = getSampleStyleSheet()
    footer_style = ParagraphStyle("Footer", parent=styles["Italic"], fontSize=8, textColor="grey")

    story = [Paragraph(profile.full_name, styles["Title"])]
    subtitle_parts = [p for p in (profile.department_name, profile.company_name) if p]
    if subtitle_parts:
        story.append(Paragraph(" — ".join(subtitle_parts), styles["Normal"]))
    story.append(
        Paragraph(f"{profile.years_experience} years of professional experience", styles["Normal"])
    )
    if profile.roles:
        story.append(Paragraph(f"Current role: {profile.roles[0].title}", styles["Normal"]))
    story.append(Spacer(1, 12))

    if profile.skills:
        story.append(Paragraph("Skills", styles["Heading2"]))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(_skill_line(s), styles["Normal"])) for s in profile.skills],
                bulletType="bullet",
            )
        )
        story.append(Spacer(1, 12))

    if profile.certifications:
        story.append(Paragraph("Certifications", styles["Heading2"]))
        story.append(
            ListFlowable(
                [
                    ListItem(Paragraph(_certification_line(c), styles["Normal"]))
                    for c in profile.certifications
                ],
                bulletType="bullet",
            )
        )
        story.append(Spacer(1, 12))

    if profile.projects:
        story.append(Paragraph("Project Experience", styles["Heading2"]))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(_project_line(p), styles["Normal"])) for p in profile.projects],
                bulletType="bullet",
            )
        )
        story.append(Spacer(1, 12))

    story.append(Paragraph(_traceability_footer(profile), footer_style))

    doc.build(story)
    return buffer.getvalue()
