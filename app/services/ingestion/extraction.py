"""Text extraction: SourceType -> ExtractedDocument.

Each extractor returns page-aware text (page_number is None where the format
has no concept of pages, e.g. CSV/JSON/TXT/MD). This is what chunking
operates on, and what page_number metadata in chunks comes from.
"""

import csv
import io
import json
from dataclasses import dataclass, field

from app.models.document import SourceType


@dataclass
class ExtractedPage:
    page_number: int | None
    text: str


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


class ExtractionError(ValueError):
    pass


def extract(source_type: SourceType, content: bytes) -> ExtractedDocument:
    extractor = _EXTRACTORS.get(source_type)
    if extractor is None:
        raise ExtractionError(f"No extractor registered for {source_type}")
    try:
        return extractor(content)
    except ExtractionError:
        raise
    except Exception as exc:  # malformed/corrupted files must fail cleanly, not crash ingestion
        raise ExtractionError(f"Failed to extract {source_type.value}: {exc}") from exc


def _extract_pdf(content: bytes) -> ExtractedDocument:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        pages = [
            ExtractedPage(page_number=i + 1, text=page.get_text()) for i, page in enumerate(doc)
        ]
    finally:
        doc.close()
    if not pages:
        raise ExtractionError("PDF has no pages.")
    return ExtractedDocument(pages=pages)


def _extract_docx(content: bytes) -> ExtractedDocument:
    import docx

    document = docx.Document(io.BytesIO(content))
    text = "\n".join(p.text for p in document.paragraphs)
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=text)])


def _extract_txt(content: bytes) -> ExtractedDocument:
    text = content.decode("utf-8", errors="replace")
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=text)])


def _extract_markdown(content: bytes) -> ExtractedDocument:
    text = content.decode("utf-8", errors="replace")
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=text)])


def _extract_csv(content: bytes) -> ExtractedDocument:
    text_io = io.StringIO(content.decode("utf-8", errors="replace"))
    reader = csv.reader(text_io)
    rows = list(reader)
    if not rows:
        raise ExtractionError("CSV has no rows.")
    header, *data_rows = rows
    lines = []
    for row in data_rows:
        pairs = ", ".join(f"{h}: {v}" for h, v in zip(header, row, strict=False))
        lines.append(pairs)
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text="\n".join(lines))])


def _extract_json(content: bytes) -> ExtractedDocument:
    data = json.loads(content.decode("utf-8"))
    text = json.dumps(data, indent=2, ensure_ascii=False)
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=text)])


def _extract_xlsx(content: bytes) -> ExtractedDocument:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    pages = []
    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header, *data_rows = rows
        lines = [f"Sheet: {sheet.title}"]
        for row in data_rows:
            pairs = ", ".join(
                f"{h}: {v}" for h, v in zip(header, row, strict=False) if h is not None
            )
            if pairs:
                lines.append(pairs)
        pages.append(ExtractedPage(page_number=sheet_index, text="\n".join(lines)))
    if not pages:
        raise ExtractionError("Workbook has no non-empty sheets.")
    return ExtractedDocument(pages=pages)


_EXTRACTORS = {
    SourceType.PDF: _extract_pdf,
    SourceType.DOCX: _extract_docx,
    SourceType.TXT: _extract_txt,
    SourceType.MARKDOWN: _extract_markdown,
    SourceType.CSV: _extract_csv,
    SourceType.JSON: _extract_json,
    SourceType.XLSX: _extract_xlsx,
}
