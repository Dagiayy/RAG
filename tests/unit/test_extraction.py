import io
import json

from app.models.document import SourceType
from app.services.ingestion.extraction import extract


def test_extract_txt():
    doc = extract(SourceType.TXT, b"hello world")
    assert doc.full_text == "hello world"


def test_extract_markdown():
    doc = extract(SourceType.MARKDOWN, b"# Title\n\nBody text")
    assert "Title" in doc.full_text
    assert "Body text" in doc.full_text


def test_extract_csv():
    content = b"name,role\nAlice,Engineer\nBob,Manager"
    doc = extract(SourceType.CSV, content)
    assert "name: Alice" in doc.full_text
    assert "role: Engineer" in doc.full_text


def test_extract_json():
    content = json.dumps({"key": "value", "nested": {"a": 1}}).encode()
    doc = extract(SourceType.JSON, content)
    assert "key" in doc.full_text
    assert "value" in doc.full_text


def test_extract_docx():
    import docx

    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(buffer)

    doc = extract(SourceType.DOCX, buffer.getvalue())
    assert "First paragraph." in doc.full_text
    assert "Second paragraph." in doc.full_text


def test_extract_pdf():
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Hello from PDF")
    content = pdf.tobytes()
    pdf.close()

    doc = extract(SourceType.PDF, content)
    assert len(doc.pages) == 1
    assert doc.pages[0].page_number == 1
    assert "Hello from PDF" in doc.full_text


def test_extract_xlsx():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "People"
    sheet.append(["name", "role"])
    sheet.append(["Alice", "Engineer"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    doc = extract(SourceType.XLSX, buffer.getvalue())
    assert "name: Alice" in doc.full_text
    assert "role: Engineer" in doc.full_text
